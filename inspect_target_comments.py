import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
path = r"D:\New folder\G6 Scoring System Final 2025-2026-MLCIS -.xlsm"
book = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
for name in ("Q1", "Comment Q1", "Q4", "Comment Q4"):
    sheet=book[name]
    print('\n', name)
    for r in range(1, 17):
        values=[]
        for c in range(1, min(sheet.max_column, 80)+1):
            value=sheet.cell(r,c).value
            if value is not None:
                values.append(f'{sheet.cell(r,c).coordinate}={str(value)[:50]}')
        if values: print(' | '.join(values))
book.close()
