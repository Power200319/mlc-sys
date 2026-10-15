from pathlib import Path
from tempfile import TemporaryDirectory
import zipfile

from django.test import Client, SimpleTestCase, override_settings
from openpyxl import Workbook, load_workbook

from .mapping import MAPPING_BY_KEY
from .services import (
    Student,
    SubjectUpload,
    build_subject_uploads,
    consume_output,
    fill_merge_fields,
    populate_scoring_template,
    report_values,
    roster_matches,
)


class BatchImportTests(SimpleTestCase):
    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_process_endpoint_accepts_multiple_upload_requests_without_csrf(self):
        client = Client(enforce_csrf_checks=True)
        response = client.post("/api/process/")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"], "Choose one or more .xlsx E-gradebook files.")

    def test_roster_accepts_same_names_when_ids_are_blank(self):
        first = [Student(11, 1, "Kinh Rothavotey", "", "", "F", "III", 80, (None, None, None))]
        second = [Student(11, 1, "Kinh Rothavotey", "", "", "F", "III", 90, (None, None, None))]
        self.assertTrue(roster_matches(first, second))

    def test_report_values_use_imported_score_and_never_template_sample_score(self):
        student = Student(11, 1, "Kinh Rothavotey", "គិញ រដ្ឋាវត្តី", "0205", "F", "III", 88.75, (None, None, None))
        upload = SubjectUpload("english.xlsx", "English", "Khat Tho", MAPPING_BY_KEY["english"], "Q1", [student])
        values = report_values(0, [upload], "2025-2026", 1)
        self.assertEqual(values["F2"], "Kinh Rothavotey")
        self.assertEqual(values["M_50"], "គិញ រដ្ឋាវត្តី")
        self.assertEqual(values["F16"], "88.75")
        self.assertEqual(values["F32"], "B")
        self.assertEqual(values["F18"], "")
        self.assertEqual(values["F108"], "")

    def test_filling_word_fields_preserves_namespace_prefixes(self):
        """Word requires prefixes named by mc:Ignorable to remain unchanged."""
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.docx"
            output = Path(directory) / "student.docx"
            document_xml = b'''<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:w14="urn:test" mc:Ignorable="w14"><w:body><w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> MERGEFIELD F2 </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Sample student</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p><w:p><w:r><w:t>Fourth</w:t></w:r></w:p></w:body></w:document>'''
            with zipfile.ZipFile(template, "w") as archive:
                archive.writestr("word/document.xml", document_xml)
                archive.writestr("custom/untouched.bin", b"unchanged")

            fill_merge_fields(template, output, {"F2": "Kinh Rothavotey", "ACADEMIC_YEAR": "2025-2026"}, "Q1")

            with zipfile.ZipFile(output) as archive:
                filled_xml = archive.read("word/document.xml")
                self.assertIn(b'xmlns:mc=', filled_xml)
                self.assertIn(b'mc:Ignorable="w14"', filled_xml)
                self.assertIn(b"Kinh Rothavotey", filled_xml)
                self.assertNotIn(b"Sample student", filled_xml)
                self.assertIn(b'w:fldCharType="end"', filled_xml)
                self.assertIn(b">First</w:t>", filled_xml)
                self.assertEqual(archive.read("custom/untouched.bin"), b"unchanged")

    def test_blank_subject_detects_multiple_subject_tabs(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "all-subjects.xlsx"
            workbook = Workbook()
            settings = workbook.active
            settings.title = "SETTINGS"
            settings["F23"] = ""
            for name in ("Khmer Dictations", "SPORT"):
                sheet = workbook.create_sheet(name)
                sheet["O9"] = "Letter Grade"
                sheet["N9"] = "Total" if name == "SPORT" else None
                if name != "SPORT":
                    sheet["O9"] = "Total"
                    sheet["P9"] = "Letter Grade"
                else:
                    sheet["O9"] = "Letter Grade"
                sheet["C11"] = "Kinh Rothavotey"
                sheet["D11"] = "គិញ រដ្ឋាវត្តី"
                sheet["F11"] = "F"
                sheet["G11"] = "III"
                sheet["N11" if name == "SPORT" else "O11"] = 80
            workbook.save(path)
            workbook.close()
            uploads = build_subject_uploads(path, path.name)
            self.assertEqual([upload.mapping.key for upload in uploads], ["khmer_dictation", "sport"])

    def test_q1_batch_imports_english_and_khmer_dictation_to_different_columns(self):
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.xlsx"
            output = Path(directory) / "output.xlsx"
            workbook = Workbook()
            q1 = workbook.active
            q1.title = "Q1"
            q1["W15"] = '=IF(G15="","", "A")'
            q1["AF15"] = '=IF(P15="","", "A")'
            comments = workbook.create_sheet("Comment Q1")
            comments["A9"] = "=Q1!A15"
            for quarter in ("Q2", "Q3", "Q4"):
                sample = workbook.create_sheet(quarter)
                sample["B15"] = f"{quarter} Sample Student"
                sample["P15"] = 99
                sample["W15"] = '=IF(G15="","", "A")'
            for quarter in ("Comment Q2", "Comment Q3", "Comment Q4"):
                sample_comment = workbook.create_sheet(quarter)
                sample_comment["G9"] = "Old sample comment"
            workbook.save(template)
            workbook.close()

            roster = [Student(11, 1, "Kinh Rothavotey", "គិញ រដ្ឋាវត្តី", "0205", "F", "III", 88.75, ("Works hard.", "Respectful.", "Read more."))]
            english = SubjectUpload("english.xlsx", "English", "Khat Tho", MAPPING_BY_KEY["english"], "Q1", roster)
            khmer_roster = [Student(11, 1, "Kinh Rothavotey", "គិញ រដ្ឋាវត្តី", "0205", "F", "III", 92, ("បានល្អ", "មានវិន័យ", "អានបន្ថែម"))]
            dictation = SubjectUpload("dictation.xlsx", "Khmer Dictation", "Sok Dara", MAPPING_BY_KEY["khmer_dictation"], "Q1", khmer_roster)
            scores, comments_count = populate_scoring_template(template, output, "Q1", [english, dictation])
            result = load_workbook(output, data_only=False)
            self.assertEqual(scores, 2)
            self.assertEqual(comments_count, 6)
            self.assertEqual(result["Q1"]["G15"].value, 92)
            self.assertEqual(result["Q1"]["P15"].value, 88.75)
            self.assertEqual(result["Q1"]["AF15"].value, '=IF(P15="","", "A")')
            self.assertEqual(result["Comment Q1"]["G9"].value, "Works hard.")
            self.assertEqual(result["Comment Q1"]["J9"].value, "Khat Tho")
            self.assertEqual(result["Comment Q1"]["AT9"].value, "Sok Dara")
            self.assertEqual(result["Comment Q1"]["AQ9"].value, "បានល្អ")
            self.assertEqual(result["Comment Q1"]["AR9"].value, "មានវិន័យ")
            self.assertEqual(result["Comment Q1"]["AS9"].value, "អានបន្ថែម")
            self.assertIsNone(result["Q2"]["B15"].value)
            self.assertIsNone(result["Q3"]["P15"].value)
            self.assertIsNone(result["Q4"]["B15"].value)
            self.assertIsNone(result["Comment Q2"]["G9"].value)
            self.assertIsNone(result["Comment Q3"]["G9"].value)
            self.assertIsNone(result["Comment Q4"]["G9"].value)
            result.close()

    def test_continued_workbook_preserves_previous_quarter_and_comments(self):
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.xlsx"
            q1_output = Path(directory) / "q1.xlsx"
            q2_output = Path(directory) / "q2.xlsx"
            workbook = Workbook()
            q1 = workbook.active
            q1.title = "Q1"
            q1["P15"] = 99
            q1["AF15"] = '=IF(P15="","", "A")'
            comment_q1 = workbook.create_sheet("Comment Q1")
            comment_q1["A9"] = "=Q1!A15"
            comment_q1["G9"] = "Old comment"
            q2 = workbook.create_sheet("Q2")
            q2["P15"] = 98
            q2["AF15"] = '=IF(P15="","", "A")'
            comment_q2 = workbook.create_sheet("Comment Q2")
            comment_q2["A9"] = "=Q2!A15"
            comment_q2["G9"] = "Old comment"
            for quarter in ("Q3", "Q4"):
                workbook.create_sheet(quarter)["AF15"] = '=IF(P15="","", "A")'
            for quarter in ("Comment Q3", "Comment Q4"):
                workbook.create_sheet(quarter)
            workbook.save(template)
            workbook.close()

            q1_roster = [Student(11, 1, "Kinh Rothavotey", "គិញ រដ្ឋាវត្តី", "0205", "F", "III", 81, ("Q1 effort", None, None))]
            q1_english = SubjectUpload("q1.xlsx", "English", "Khat Tho", MAPPING_BY_KEY["english"], "Q1", q1_roster)
            populate_scoring_template(template, q1_output, "Q1", [q1_english], clear_samples=True)

            q2_roster = [Student(11, 1, "Kinh Rothavotey", "គិញ រដ្ឋាវត្តី", "0205", "F", "III", 91, ("Q2 effort", None, None))]
            q2_english = SubjectUpload("q2.xlsx", "English", "Khat Tho", MAPPING_BY_KEY["english"], "Q2", q2_roster)
            populate_scoring_template(q1_output, q2_output, "Q2", [q2_english], clear_samples=False)

            result = load_workbook(q2_output, data_only=False)
            self.assertEqual(result["Q1"]["P15"].value, 81)
            self.assertEqual(result["Comment Q1"]["G9"].value, "Q1 effort")
            self.assertEqual(result["Q2"]["P15"].value, 91)
            self.assertEqual(result["Comment Q2"]["G9"].value, "Q2 effort")
            result.close()

    def test_download_consumes_output_and_removes_server_copy(self):
        with TemporaryDirectory() as directory:
            jobs_dir = Path(directory) / "jobs"
            job_dir = jobs_dir / "safejob"
            job_dir.mkdir(parents=True)
            workbook = job_dir / "G4-G6 Scoring System Final 2024-2025-MLCIS -6A.xlsm"
            workbook.write_bytes(b"private student data")
            with override_settings(MLCIS_JOB_DIR=jobs_dir, MLCIS_JOB_RETENTION_SECONDS=600):
                content, filename = consume_output("safejob", "scoring")
            self.assertEqual(content, b"private student data")
            self.assertEqual(filename, workbook.name)
            self.assertFalse(workbook.exists())
            self.assertFalse(job_dir.exists())
