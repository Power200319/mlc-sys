"""Batch-import subject E-gradebooks into a fresh scoring-format copy."""

from __future__ import annotations

import uuid
import zipfile
import time
from copy import copy
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any
import re
import shutil
import subprocess
from xml.sax.saxutils import escape

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from .mapping import COMMENT_SHEETS, MAPPING_BY_SOURCE_SUBJECT, MULTI_SUBJECT_SHEETS, PERIOD_SHEETS, SUPPORTED_PERIODS, SubjectMapping, normalise_subject


class ProcessingError(Exception):
    pass


OUTPUT_FILENAME = "G4-G6 Scoring System Final 2024-2025-MLCIS -6A.xlsm"
REPORTS_FILENAME = "Grade-6A-Student-Report-Cards.zip"


@dataclass
class Student:
    source_row: int
    number: int | None
    english_name: str
    khmer_name: str
    student_id: str
    sex: str
    group: str
    total: float | None
    comments: tuple[str | None, str | None, str | None]


@dataclass
class SubjectUpload:
    filename: str
    subject_name: str
    teacher_name: str
    mapping: SubjectMapping
    source_period: str | None
    students: list[Student]


def is_xlsx_upload(upload) -> bool:
    upload.seek(0)
    valid = upload.name.lower().endswith(".xlsx") and zipfile.is_zipfile(upload)
    upload.seek(0)
    return valid


def is_scoring_workbook_upload(upload) -> bool:
    """A continued scoring workbook must remain an Excel macro workbook."""
    if upload is None:
        return False
    upload.seek(0)
    valid = upload.name.lower().endswith(".xlsm") and zipfile.is_zipfile(upload)
    upload.seek(0)
    return valid


def configured_templates() -> dict[str, Path]:
    return {
        "scoring": Path(settings.MLCIS_SCORING_TEMPLATE),
        "quarter_report": Path(settings.MLCIS_QUARTER_REPORT_TEMPLATE),
        "semester_report": Path(settings.MLCIS_SEMESTER_REPORT_TEMPLATE),
        "end_year_report": Path(settings.MLCIS_END_YEAR_REPORT_TEMPLATE),
    }


def purge_expired_jobs() -> None:
    """Remove unclaimed output files after the short privacy retention window."""
    jobs_dir = Path(settings.MLCIS_JOB_DIR)
    if not jobs_dir.exists():
        return
    cutoff = time.time() - settings.MLCIS_JOB_RETENTION_SECONDS
    for job_dir in jobs_dir.iterdir():
        if job_dir.is_dir() and job_dir.stat().st_mtime < cutoff:
            shutil.rmtree(job_dir, ignore_errors=True)


def _source_setting(workbook_path: Path, cell: str) -> str:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "SETTINGS" not in workbook.sheetnames:
            raise ProcessingError("An uploaded file does not contain a SETTINGS sheet.")
        return str(workbook["SETTINGS"][cell].value or "").strip()
    finally:
        workbook.close()


def source_period(workbook_path: Path) -> str | None:
    value = _source_setting(workbook_path, "F14")
    normalised = value.upper().replace("QUARTER", "Q").replace(" ", "")
    return normalised if normalised in SUPPORTED_PERIODS else None


def source_subject(workbook_path: Path) -> str:
    return _source_setting(workbook_path, "F23")


def source_teacher_name(workbook_path: Path) -> str:
    """Teacher's Name is stored in SETTINGS!F20 in the supplied E-gradebook."""
    return _source_setting(workbook_path, "F20")


def read_subject_students(workbook_path: Path) -> list[Student]:
    """Read displayed values from the common E-gradebook Scoring System sheet."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Scoring System" not in workbook.sheetnames:
            raise ProcessingError("An uploaded file does not contain a Scoring System sheet.")
        sheet = workbook["Scoring System"]
        if str(sheet["AM9"].value or "").strip().casefold() != "total":
            raise ProcessingError("An uploaded file has no verified Scoring System Total column (AM9).")
        students: list[Student] = []
        blank_streak = 0
        for row in range(11, sheet.max_row + 1):
            name = sheet[f"C{row}"].value
            if not name:
                blank_streak += 1
                if students and blank_streak >= 3:
                    break
                continue
            blank_streak = 0
            total = sheet[f"AM{row}"].value
            students.append(Student(
                source_row=row,
                number=sheet[f"B{row}"].value if isinstance(sheet[f"B{row}"].value, int) else None,
                english_name=str(name).strip(),
                khmer_name=str(sheet[f"D{row}"].value or "").strip(),
                student_id=str(sheet[f"E{row}"].value or "").strip(),
                sex=str(sheet[f"F{row}"].value or "").strip(),
                group=str(sheet[f"G{row}"].value or "").strip(),
                total=float(total) if isinstance(total, (int, float)) else None,
                comments=tuple(str(sheet[f"{column}{row}"].value or "").strip() or None for column in ("AP", "AQ", "AR")),
            ))
        if not students:
            raise ProcessingError("No student rows with displayed names were found in Scoring System!C11:C.")
        return students
    finally:
        workbook.close()


def _first_nonempty(values) -> str:
    return next((str(value).strip() for value in values if value not in (None, "")), "")


def read_multi_subject_sheet(workbook_path: Path, sheet_name: str) -> tuple[list[Student], str]:
    """Read one actual subject worksheet from an all-subject gradebook.

    These sheets share the Gradebook row structure: names start in row 11,
    final Total is immediately before Letter Grade, and comments/teacher name
    follow the Rank column.  We discover the columns from headers so a sheet
    with different numbers of homework/tests still imports correctly.
    """
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        limit = min(sheet.max_column, 100)
        total_column = next((column for column in range(1, limit) if str(sheet.cell(9, column).value or "").strip().casefold() == "total" and "letter" in str(sheet.cell(9, column + 1).value or "").casefold()), None)
        # Health Education in the supplied multi-subject workbook has blank
        # row-9 headers but uses the same known final Total position.
        if total_column is None and sheet_name == "Health Education ":
            total_column = 14
        if total_column is None:
            raise ProcessingError(f"{sheet_name}: could not locate final Total beside Letter Grade.")
        teacher_column = next((column for column in range(1, limit + 1) if "teacher" in str(sheet.cell(9, column).value or "").casefold()), None)
        comment_columns = (teacher_column - 3, teacher_column - 2, teacher_column - 1) if teacher_column and teacher_column > 3 else None
        students: list[Student] = []
        blank_streak = 0
        teacher_names: list[str] = []
        for row in range(11, sheet.max_row + 1):
            name = sheet[f"C{row}"].value
            if not name:
                blank_streak += 1
                if students and blank_streak >= 3:
                    break
                continue
            blank_streak = 0
            total = sheet.cell(row, total_column).value
            if teacher_column:
                teacher_names.append(str(sheet.cell(row, teacher_column).value or "").strip())
            comments = tuple(
                str(sheet.cell(row, column).value or "").strip() or None
                for column in comment_columns
            ) if comment_columns else (None, None, None)
            students.append(Student(
                source_row=row,
                number=sheet[f"B{row}"].value if isinstance(sheet[f"B{row}"].value, int) else None,
                english_name=str(name).strip(),
                khmer_name=str(sheet[f"D{row}"].value or "").strip(),
                student_id=str(sheet[f"E{row}"].value or "").strip(),
                sex=str(sheet[f"F{row}"].value or "").strip(),
                group=str(sheet[f"G{row}"].value or "").strip(),
                total=float(total) if isinstance(total, (int, float)) else None,
                comments=comments,
            ))
        if not students:
            raise ProcessingError(f"{sheet_name}: no student rows were found.")
        return students, _first_nonempty(teacher_names)
    finally:
        workbook.close()


def ensure_scoring_template_is_usable(path: Path) -> None:
    if not path.exists():
        raise ProcessingError("The scoring template file is not configured or cannot be found.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        required = set(PERIOD_SHEETS.values()) | set(COMMENT_SHEETS.values())
        if not required.issubset(workbook.sheetnames):
            raise ProcessingError("The scoring template is missing a required Quarter or Comment Quarter sheet.")
        if not workbook.vba_archive:
            raise ProcessingError("The scoring template VBA project could not be read.")
    finally:
        workbook.close()


def roster_matches(reference: list[Student], candidate: list[Student]) -> bool:
    if len(reference) != len(candidate):
        return False
    for left, right in zip(reference, candidate):
        if left.student_id and right.student_id:
            if left.student_id != right.student_id:
                return False
        elif left.english_name.casefold() != right.english_name.casefold():
            return False
    return True


def clear_sample_constants(sheet, first_data_row: int, last_column: int) -> None:
    """Remove sample data without deleting the format's formulas or styling."""
    for row in range(first_data_row, sheet.max_row + 1):
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.value = None


def clear_all_quarters(workbook) -> None:
    """An output contains data only for quarters imported in this job."""
    for quarter in PERIOD_SHEETS.values():
        # A is a row-number formula and W onward contains score formulas.
        # Clearing constants in A:V removes all example roster/score values.
        clear_sample_constants(workbook[quarter], 15, 22)
    for comment_quarter in COMMENT_SHEETS.values():
        # Keep formula links in A:F but remove all sample comments and names.
        clear_sample_constants(workbook[comment_quarter], 9, workbook[comment_quarter].max_column)


def extend_quarter_rows(quarter_sheet, required_students: int) -> None:
    """Extend the final student row when the supplied format is too short."""
    first_row = 15
    existing_students = quarter_sheet.max_row - first_row + 1
    while existing_students < required_students:
        source_row = first_row + existing_students - 1
        destination_row = source_row + 1
        quarter_sheet.row_dimensions[destination_row].height = quarter_sheet.row_dimensions[source_row].height
        for column in range(1, quarter_sheet.max_column + 1):
            source = quarter_sheet.cell(source_row, column)
            target = quarter_sheet.cell(destination_row, column)
            if isinstance(source.value, str) and source.value.startswith("="):
                target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
            else:
                target.value = source.value
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.alignment:
                target.alignment = copy(source.alignment)
            if source.font:
                target.font = copy(source.font)
            if source.fill:
                target.fill = copy(source.fill)
            if source.border:
                target.border = copy(source.border)
            if source.protection:
                target.protection = copy(source.protection)
        existing_students += 1


def populate_scoring_template(
    template_path: Path,
    output_path: Path,
    period: str,
    uploads: list[SubjectUpload],
    *,
    clear_samples: bool = True,
) -> tuple[int, int]:
    workbook = load_workbook(template_path, keep_vba=template_path.suffix.lower() == ".xlsm", keep_links=True)
    try:
        quarter_sheet = workbook[PERIOD_SHEETS[period]]
        comment_sheet = workbook[COMMENT_SHEETS[period]]
        roster = uploads[0].students
        first_target_row, first_comment_row = 15, 9
        comment_capacity = comment_sheet.max_row - first_comment_row + 1
        if len(roster) > comment_capacity:
            raise ProcessingError(f"The {period} Comment format has room for {comment_capacity} students, but an upload contains {len(roster)}.")
        extend_quarter_rows(quarter_sheet, len(roster))

        # A new workbook starts from a clean format. A continued workbook is
        # already a real teacher file, so never clear its other quarters,
        # comments, or manually entered Exam 1 / Exam 2 values.
        if clear_samples:
            clear_all_quarters(workbook)
        for index, student in enumerate(roster):
            row = first_target_row + index
            quarter_sheet[f"A{row}"] = student.number or index + 1
            quarter_sheet[f"B{row}"] = student.english_name
            quarter_sheet[f"C{row}"] = student.khmer_name
            quarter_sheet[f"D{row}"] = student.sex
            quarter_sheet[f"E{row}"] = student.student_id
            quarter_sheet[f"F{row}"] = student.group

        transferred_scores = transferred_comments = 0
        for upload in uploads:
            for index, student in enumerate(upload.students):
                row = first_target_row + index
                if student.total is not None:
                    quarter_sheet[f"{upload.mapping.target_score_column}{row}"] = student.total
                    transferred_scores += 1
                if upload.mapping.comment_columns:
                    for column, value in zip(upload.mapping.comment_columns, student.comments):
                        if value:
                            comment_sheet[f"{column}{first_comment_row + index}"] = value
                            transferred_comments += 1
                if upload.mapping.teacher_name_column and upload.teacher_name:
                    comment_sheet[f"{upload.mapping.teacher_name_column}{first_comment_row + index}"] = upload.teacher_name

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(output_path)
        return transferred_scores, transferred_comments
    finally:
        workbook.close()


REPORT_SCORE_KEYS = (
    "khmer_dictation", "khmer_essays", "khmer_reading", "khmer_public_speaking",
    "khmer_math", "khmer_social_studies", "khmer_science", "sport", "health_education",
    "english", "english_grammar", "english_math", "english_science", "chinese", "robothink", "computer",
)
REPORT_CREDITS = (2, 2, 3, 1, 5, 3, 2, 1, 1, 5, 1, 5, 3, 2, 2, 2)


def letter_grade(score: float | None) -> str:
    if score is None:
        return ""
    if score >= 90: return "A"
    if score >= 80: return "B"
    if score >= 70: return "C"
    if score >= 60: return "D"
    if score >= 50: return "E"
    return "F"


def grade_point(grade: str) -> float:
    return {"A": 4, "B": 3, "C": 2, "D": 1.5, "E": 1, "F": 0}.get(grade, 0)


def display_score(score: float | None) -> str:
    return f"{score:.2f}" if score is not None else ""


def source_academic_year(path: Path) -> str:
    return _source_setting(path, "F11")


def report_values(student_index: int, uploads: list[SubjectUpload], academic_year: str, class_size: int) -> dict[str, str]:
    roster_student = uploads[0].students[student_index]
    uploads_by_key = {item.mapping.key: item for item in uploads}
    scores = [
        uploads_by_key[key].students[student_index].total if key in uploads_by_key else None
        for key in REPORT_SCORE_KEYS
    ]
    grades = [letter_grade(score) for score in scores]
    khmer_scores, english_scores = scores[:9], scores[9:]
    present_scores = [score for score in scores if score is not None]
    khmer_present = [score for score in khmer_scores if score is not None]
    english_present = [score for score in english_scores if score is not None]
    overall = sum(present_scores) / len(present_scores) if present_scores else None
    ranking = []
    for row in range(class_size):
        row_scores = [item.students[row].total for item in uploads if item.students[row].total is not None]
        if row_scores:
            ranking.append((sum(row_scores) / len(row_scores), row))
    ranking.sort(key=lambda value: value[0], reverse=True)
    rank = next((position + 1 for position, (_, row) in enumerate(ranking) if row == student_index), "")
    earned_credits = sum(credit for score, credit in zip(scores, REPORT_CREDITS) if score is not None)
    gpa = sum(grade_point(grade) for grade in grades if grade) / len([grade for grade in grades if grade]) if any(grades) else None
    values = {
        "F2": roster_student.english_name,
        "F4": roster_student.sex,
        "F5": roster_student.student_id,
        "M_50": roster_student.khmer_name,
        "ACADEMIC_YEAR": academic_year,
        "F90": display_score(sum(khmer_present) / len(khmer_present) if khmer_present else None),
        "F91": letter_grade(sum(khmer_present) / len(khmer_present) if khmer_present else None),
        "F87": str(sum(credit for score, credit in zip(khmer_scores, REPORT_CREDITS[:9]) if score is not None)) if khmer_present else "",
        "F96": display_score(sum(english_present) / len(english_present) if english_present else None),
        "F97": letter_grade(sum(english_present) / len(english_present) if english_present else None),
        "F93": str(sum(credit for score, credit in zip(english_scores, REPORT_CREDITS[9:]) if score is not None)) if english_present else "",
        "F99": str(earned_credits) if present_scores else "",
        "F102": display_score(overall),
        "F103": letter_grade(overall),
        "F104": f"{gpa:.2f}" if gpa is not None else "",
        "F105": str(rank),
        "F107": str(class_size),
        # Attendance is not in the imported score workbooks. Blank fields are
        # safer than retaining the template student's attendance numbers.
        "F108": "", "F109": "", "F110": "", "F111": "", "F112": "", "F113": "",
        "Total_School_Day": "",
    }
    for number, (score, grade, credit) in enumerate(zip(scores, grades, REPORT_CREDITS), start=7):
        values[f"F{number}"] = display_score(score)
        values[f"F{number + 16}"] = grade
        values[f"F{number + 32}"] = str(credit) if score is not None else ""
    return values


def fill_merge_fields(template_path: Path, output_path: Path, values: dict[str, str], period: str) -> None:
    """Fill cached merge-field results without serializing/rebuilding Word XML.

    Re-serializing XML changes namespace prefixes used by mc:Ignorable and can
    make Microsoft Word report the document as corrupt. This byte-preserving
    patch changes only text inside the existing field-result nodes.
    """
    quarter_names = {"Q1": "First", "Q2": "Second", "Q3": "Third", "Q4": "Fourth"}
    field_pattern = re.compile(
        rb'(<w:fldChar\b[^>]*w:fldCharType="begin"[^>]*/>.*?<w:instrText\b[^>]*>(?P<instruction>.*?)</w:instrText>.*?<w:fldChar\b[^>]*w:fldCharType="separate"[^>]*/>)(?P<result>.*?)(<w:fldChar\b[^>]*w:fldCharType="end"[^>]*/>)',
        re.DOTALL,
    )
    text_pattern = re.compile(rb'(<w:t\b[^>]*>)(.*?)(</w:t>)', re.DOTALL)

    def replace_field(match: re.Match[bytes]) -> bytes:
        instruction = match.group("instruction").decode("utf-8", errors="ignore")
        field = re.search(r'MERGEFIELD\s+"?([^"\\ ]+)', instruction)
        if not field:
            return match.group(0)
        replacement = escape(values.get(field.group(1), "")).encode("utf-8")
        result, changed = text_pattern.subn(lambda text: text.group(1) + replacement + text.group(3), match.group("result"), count=1)
        # Groups are: field prefix, instruction, existing result, field-end.
        # Keep the end marker. Re-appending group 3 would duplicate the old
        # sample value and remove the field terminator, causing Word to print
        # MERGEFIELD instructions in the exported PDF.
        return match.group(1) + (result if changed else match.group("result")) + match.group(4)

    with zipfile.ZipFile(template_path) as source, zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as destination:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "word/document.xml":
                content = field_pattern.sub(replace_field, content)
                content = content.replace(b">Fourth</w:t>", f">{quarter_names[period]}</w:t>".encode("utf-8"))
                content = content.replace(b">2025-2026</w:t>", f">{escape(values.get('ACADEMIC_YEAR', ''))}</w:t>".encode("utf-8"))
            destination.writestr(info, content)


def convert_word_reports_to_pdf(reports_dir: Path) -> None:
    script = Path(__file__).with_name("word_to_pdf.ps1")
    command = ["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script), str(reports_dir)]
    completed = subprocess.run(command, capture_output=True, text=True, timeout=180)
    expected = list(reports_dir.glob("*.docx"))
    if completed.returncode != 0 or any(not document.with_suffix(".pdf").exists() for document in expected):
        detail = (completed.stderr or completed.stdout or "Microsoft Word could not export the PDF.").strip()
        raise ProcessingError(f"PDF export failed: {detail[:300]}")


def safe_filename(value: str) -> str:
    return re.sub(r"[<>:\\/?*\"|]", "_", value).strip(" .") or "student"


def generate_report_cards(template_path: Path, job_dir: Path, uploads: list[SubjectUpload], academic_year: str, period: str) -> tuple[Path, int]:
    if not template_path.exists():
        raise ProcessingError("The report-card Word template file could not be found.")
    reports_dir = job_dir / "student-report-cards"
    reports_dir.mkdir(exist_ok=True)
    for index, student in enumerate(uploads[0].students, start=1):
        stem = f"{index:02d}-{safe_filename(student.english_name)}"
        docx_path = reports_dir / f"{stem}.docx"
        values = report_values(index - 1, uploads, academic_year, len(uploads[0].students))
        fill_merge_fields(template_path, docx_path, values, period)
    convert_word_reports_to_pdf(reports_dir)
    for docx_path in reports_dir.glob("*.docx"):
        docx_path.unlink()
    archive_base = job_dir / "Grade-6A-Student-Report-Cards"
    archive = Path(shutil.make_archive(str(archive_base), "zip", reports_dir))
    return archive, len(uploads[0].students)


def build_subject_uploads(path: Path, filename: str) -> list[SubjectUpload]:
    subject_name = source_subject(path)
    mapping = MAPPING_BY_SOURCE_SUBJECT.get(normalise_subject(subject_name))
    if mapping is not None:
        return [SubjectUpload(
            filename,
            subject_name,
            source_teacher_name(path),
            mapping,
            source_period(path),
            read_subject_students(path),
        )]

    # SETTINGS!F23 is blank in all-subject files. Detect each known subject
    # tab, allowing one uploaded file to contribute 3, 4, 5, or all subjects.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        detected = [(sheet_name, subject_mapping) for sheet_name, subject_mapping in MULTI_SUBJECT_SHEETS.items() if sheet_name in workbook.sheetnames]
    finally:
        workbook.close()
    if not detected:
        raise ProcessingError(f"{filename}: Subject is blank in SETTINGS!F23 and no supported subject sheets were found.")
    period = source_period(path)
    setting_teacher = source_teacher_name(path)
    results: list[SubjectUpload] = []
    for sheet_name, subject_mapping in detected:
        students, sheet_teacher = read_multi_subject_sheet(path, sheet_name)
        results.append(SubjectUpload(filename, sheet_name, sheet_teacher or setting_teacher, subject_mapping, period, students))
    return results


PRINT_STAGES = {
    "P1": {"period": "Q1", "label": "Print 1: Q1 and Teacher Comment", "report": True, "template": "quarter_report"},
    "P2": {"period": "Q2", "label": "Print 2: Q2, Exam 1, and Final SEM1", "report": False},
    "P3": {"period": "Q3", "label": "Print 3: Q3 and Teacher Comment", "report": True, "template": "quarter_report"},
    "P4": {"period": "Q4", "label": "Print 4: Q4, Exam 2, and Final SEM2", "report": False},
}


def process_gradebooks(uploads, period: str, scoring_workbook=None, print_stage: str = "") -> dict[str, Any]:
    purge_expired_jobs()
    if period not in SUPPORTED_PERIODS:
        raise ProcessingError("Choose Quarter 1, 2, 3, or 4.")
    if not uploads:
        raise ProcessingError("Choose one or more .xlsx E-gradebook files.")
    if any(not is_xlsx_upload(upload) for upload in uploads):
        raise ProcessingError("Every selected file must be a valid .xlsx E-gradebook.")
    stage = PRINT_STAGES.get(print_stage)
    if print_stage and stage is None:
        raise ProcessingError("Choose a valid print package.")
    if stage and stage["period"] != period:
        raise ProcessingError(f"{stage['label']} must import into {stage['period']}.")
    if stage and stage["period"] in {"Q2", "Q3", "Q4"} and scoring_workbook is None:
        raise ProcessingError(f"{stage['label']} requires the scoring .xlsm downloaded from the previous print package.")
    if scoring_workbook is not None and not is_scoring_workbook_upload(scoring_workbook):
        raise ProcessingError("The continued scoring workbook must be the downloaded .xlsm file.")
    templates = configured_templates()
    ensure_scoring_template_is_usable(templates["scoring"])

    job_id = uuid.uuid4().hex
    job_dir = Path(settings.MLCIS_JOB_DIR) / job_id
    job_dir.mkdir(parents=True, exist_ok=False)
    source_template = templates["scoring"]
    if scoring_workbook is not None:
        source_template = job_dir / "continued-scoring.xlsm"
        with source_template.open("wb+") as destination:
            for chunk in scoring_workbook.chunks():
                destination.write(chunk)
        try:
            ensure_scoring_template_is_usable(source_template)
        except Exception:
            source_template.unlink(missing_ok=True)
            raise
    subject_uploads: list[SubjectUpload] = []
    uploads_by_subject: dict[str, SubjectUpload] = {}
    replacements: list[str] = []
    academic_year = ""
    for number, upload in enumerate(uploads, start=1):
        source_path = job_dir / f"source-{number}.xlsx"
        with source_path.open("wb+") as destination:
            for chunk in upload.chunks():
                destination.write(chunk)
        try:
            if not academic_year:
                academic_year = source_academic_year(source_path)
            for subject_upload in build_subject_uploads(source_path, upload.name):
                if subject_upload.source_period and subject_upload.source_period != period:
                    raise ProcessingError(
                        f"{upload.name} / {subject_upload.mapping.label} is set to {subject_upload.source_period}; "
                        f"it cannot be imported into {period}."
                    )
                if subject_uploads and not roster_matches(subject_uploads[0].students, subject_upload.students):
                    raise ProcessingError(
                        f"{upload.name} / {subject_upload.mapping.label}: student roster does not match "
                        "the first detected subject. Use files for the same class only."
                    )
                previous = uploads_by_subject.get(subject_upload.mapping.key)
                if previous is not None:
                    subject_uploads[subject_uploads.index(previous)] = subject_upload
                    replacements.append(
                        f"{subject_upload.mapping.label}: {upload.name} replaced the earlier file {previous.filename}."
                    )
                else:
                    subject_uploads.append(subject_upload)
                uploads_by_subject[subject_upload.mapping.key] = subject_upload
        finally:
            # Parsed values now live only in memory; never retain uploaded files.
            source_path.unlink(missing_ok=True)

    scoring_output = job_dir / OUTPUT_FILENAME
    try:
        transferred_scores, transferred_comments = populate_scoring_template(
            source_template,
            scoring_output,
            period,
            subject_uploads,
            clear_samples=scoring_workbook is None,
        )
    finally:
        if scoring_workbook is not None:
            source_template.unlink(missing_ok=True)
    with zipfile.ZipFile(scoring_output) as archive:
        vba_preserved = "xl/vbaProject.bin" in archive.namelist()
    if not vba_preserved:
        raise ProcessingError("Output validation failed: the copied scoring workbook has no VBA project.")

    first = subject_uploads[0]
    subjects_without_cached_totals = [
        item.mapping.label
        for item in subject_uploads
        if not any(student.total is not None for student in item.students)
    ]
    warnings = [
        ("The format workbook is used only as an output template. Sample data in all Q sheets is cleared."
         if scoring_workbook is None else
         "The uploaded scoring workbook is continued. Existing quarters, comments, and manual exam values are preserved."),
        f"Selected destination: {PERIOD_SHEETS[period]}; imported subjects: {', '.join(item.mapping.label for item in subject_uploads)}.",
        "All uploaded subject files must have the same student roster. Matching uses student ID when present, otherwise exact English name and order.",
        "Only cached numeric Total values are imported. Missing cached totals remain blank.",
            "Achievement, Attitude, Areas to improve, and Teacher's Name are imported into the matching Comment Q subject group when source comments are present.",
        "Existing #REF! formulas and external links in the format are preserved and not repaired.",
    ]
    if subjects_without_cached_totals:
        warnings.append(
            "No cached Total values were found for: "
            + ", ".join(subjects_without_cached_totals)
            + ". Open those source sheets in Excel, recalculate, save, then upload again."
        )
    warnings.extend(replacements)
    try:
        if stage and not stage["report"]:
            exam_label = "Exam 1" if period == "Q2" else "Exam 2"
            warnings.append(
                f"Enter {exam_label} manually in the downloaded scoring workbook, then recalculate and save it in Excel. "
                f"Print the existing {'FinalSem1 reportcard' if period == 'Q2' else 'FinalSem2'} sheet from that workbook."
            )
            raise ProcessingError("Semester PDFs are not generated before the manual exam is entered.")
        reports_archive, report_count = generate_report_cards(
            templates[stage["template"]] if stage else templates["quarter_report"],
            job_dir,
            subject_uploads,
            academic_year,
            period,
        )
        report_output = {
            "available": True,
            "filename": reports_archive.name,
            "count": report_count,
            "download": f"/api/download/{job_id}/report_cards/",
        }
    except (ProcessingError, OSError, subprocess.TimeoutExpired) as error:
        warnings.append(str(error))
        report_output = {"available": False, "reason": "Student PDF reports could not be created."}
    return {
        "job_id": job_id,
        "print_stage": print_stage or None,
        "print_label": stage["label"] if stage else None,
        "period": period,
        "student_count": len(first.students),
        "populated_students": len(first.students),
        "subjects": [{"name": item.mapping.label, "source_subject": item.subject_name, "source_period": item.source_period} for item in subject_uploads],
        "warnings": warnings,
        "outputs": {
            "scoring": {"filename": scoring_output.name, "available": True, "download": f"/api/download/{job_id}/scoring/", "vba_preserved": vba_preserved, "fields_transferred": transferred_scores, "comments_transferred": transferred_comments},
            "report_cards": report_output,
        },
    }


def output_path(job_id: str, kind: str) -> Path:
    allowed = {"scoring": OUTPUT_FILENAME, "report_cards": REPORTS_FILENAME}
    if kind not in allowed or not job_id.isalnum():
        raise ProcessingError("Requested output was not found.")
    path = Path(settings.MLCIS_JOB_DIR) / job_id / allowed[kind]
    if not path.exists():
        raise ProcessingError("Requested output was not found.")
    return path


def consume_output(job_id: str, kind: str) -> tuple[bytes, str]:
    """Read a one-time download into memory, then erase its server copy."""
    purge_expired_jobs()
    path = output_path(job_id, kind)
    content = path.read_bytes()
    filename = path.name
    path.unlink(missing_ok=True)
    if kind == "report_cards":
        shutil.rmtree(path.parent / "student-report-cards", ignore_errors=True)
    job_dir = path.parent
    if not any(job_dir.iterdir()):
        job_dir.rmdir()
    return content, filename
