import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
path = r"D:\New folder\G6 Scoring System Final 2025-2026-MLCIS -.xlsm"
workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
for sheet_name in ["Q1", "Comment Q1", "Q4", "Comment Q4"]:
    sheet = workbook[sheet_name]
    print(f"\n--- {sheet_name}")
    for row in range(1, 17):
        values = [
            f"{sheet.cell(row, column).coordinate}={sheet.cell(row, column).value}"
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(row, column).value is not None
        ]
        if values:
            print(" | ".join(values))
