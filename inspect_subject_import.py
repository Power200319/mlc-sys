import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

for path, sheets in [
    (r"D:\New folder\Telegram Desktop\E-Gradebook_6A_English Q4.xlsx", ["SETTINGS", "Scoring System", "COMMENTS"]),
    (r"D:\New folder\G6 Scoring System Final 2025-2026-MLCIS -.xlsm", ["Q1", "Comment Q1", "Q4", "Comment Q4"]),
]:
    print("\nFILE", path)
    book = load_workbook(path, read_only=True, data_only=False, keep_vba=path.endswith(".xlsm"))
    for name in sheets:
        sheet = book[name]
        print("\nSHEET", name)
        for row in range(1, min(sheet.max_row, 18) + 1):
            cells=[]
            for col in range(1, min(sheet.max_column, 100) + 1):
                value=sheet.cell(row,col).value
                if value is not None:
                    cells.append(f"{sheet.cell(row,col).coordinate}={str(value)[:60]}")
            if cells: print(" | ".join(cells))
    book.close()
